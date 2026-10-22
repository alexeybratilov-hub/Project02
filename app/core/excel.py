"""Общие средства импорта/экспорта Excel (IMPORT_EXPORT.md, T2-05/T2-06).

Экспорт: .xlsx через openpyxl; при выгрузке в CSV — UTF-8 с BOM (§6);
защита от формульных инъекций: значения, начинающиеся с = + - @,
экранируются ведущим апострофом (§6).
Импорт: чтение первого листа, заголовки из первой строки, нормализация
(регистр/пробелы), альтернативные названия столбцов (§2).
"""

from __future__ import annotations

import csv
import io
from pathlib import Path

from openpyxl import Workbook, load_workbook

INJECTION_PREFIXES = ("=", "+", "-", "@")


def escape_cell(value):
    """Экранирование ячеек-формулок (IMPORT_EXPORT §6)."""
    if isinstance(value, str) and value.startswith(INJECTION_PREFIXES):
        return "'" + value
    return value


def normalize_header(text) -> str:
    if text is None:
        return ""
    return " ".join(str(text).strip().lower().split())


def read_xlsx_rows(path: str | Path) -> tuple[list[str], list[list]]:
    """Вернуть (заголовки, строки) первого листа книги."""
    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header_row = next(rows, None)
        if header_row is None:
            return [], []
        headers = [str(h).strip() if h is not None else "" for h in header_row]
        data = [list(r) for r in rows]
        return headers, data
    finally:
        wb.close()


def write_xlsx(headers: list[str], rows: list[list], path: str | Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(list(headers))
    for row in rows:
        ws.append([escape_cell(v) for v in row])
    wb.save(str(path))
    wb.close()


def xlsx_bytes(headers: list[str], rows: list[list]) -> bytes:
    """Сформировать .xlsx в памяти (для ответа HTTP / скачивания шаблона)."""
    buf = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.append(list(headers))
    for row in rows:
        ws.append([escape_cell(v) for v in row])
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def csv_bytes_utf8_bom(headers: list[str], rows: list[list]) -> bytes:
    """CSV в UTF-8 с BOM для корректного открытия в Excel (IMPORT_EXPORT §6)."""
    buf = io.StringIO(newline="")
    writer = csv.writer(buf)
    writer.writerow([escape_cell(h) for h in headers])
    for row in rows:
        writer.writerow([escape_cell(v) for v in row])
    return b"\xef\xbb\xbf" + buf.getvalue().encode("utf-8")  # UTF-8 BOM
