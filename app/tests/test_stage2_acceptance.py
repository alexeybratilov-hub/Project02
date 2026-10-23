"""Сквозные HTTP-сценарии этапа 2 (TESTING.md §3, сценарии 1–5, 12, 14).

Проверяются пользовательские пути через страницы и API модуля
«Юридические лица»: создание, второй записи, импорт, ошибки импорта,
отмена импорта без повреждения данных, предупреждение о дубликате, экспорт.
"""

from __future__ import annotations

import io

from openpyxl import Workbook


def _xlsx(headers, rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


HEADERS = ["Полное наименование", "Сокращённое наименование", "ИНН", "КПП",
           "ОГРН", "Юридический адрес", "Фактический адрес", "Дополнительные сведения"]


class TestScenariosViaHttp:
    def test_scenario_1_create_first_legal_entity(self, client):
        """Сценарий 1: создание юридического лица вручную."""
        page = client.get("/legal-entities/new")
        assert page.status_code == 200
        resp = client.post(
            "/legal-entities/new",
            data={"full_name": "ООО Сквозное", "inn": "7701234567"},
            follow_redirects=False,
        )
        assert resp.status_code == 303
        card = client.get(resp.headers["location"])
        assert card.status_code == 200 and "ООО Сквозное" in card.text
        api = client.get("/legal-entities/api/list").json()
        assert api["total"] == 1

    def test_scenario_2_create_second_legal_entity(self, client):
        """Сценарий 2: создание второго юридического лица (разделение данных FR-01)."""
        client.post("/legal-entities/new", data={"full_name": "ООО Первое"})
        client.post("/legal-entities/new", data={"full_name": "АО Второе"})
        api = client.get("/legal-entities/api/list").json()
        assert api["total"] == 2
        names = {i["full_name"] for i in api["items"]}
        assert names == {"ООО Первое", "АО Второе"}

    def test_validation_error_shown_in_russian(self, client):
        resp = client.post("/legal-entities/new", data={"full_name": "ООО Плохо", "inn": "123"})
        assert resp.status_code == 422
        assert "ИНН должен содержать 10 или 12 цифр" in resp.text

    def test_duplicate_conflict_on_create(self, client):
        client.post("/legal-entities/new",
                    data={"full_name": "ООО Оригинал", "inn": "7701234567", "kpp": "770101001"})
        resp = client.post("/legal-entities/new",
                           data={"full_name": "ООО Клон", "inn": "7701234567", "kpp": "770101001"})
        assert resp.status_code == 409
        assert "уже зарегистрировано" in resp.text

    def test_scenario_3_import_from_excel(self, client):
        """Сценарий 3: импорт юридических лиц из Excel (предпросмотр → подтверждение)."""
        content = _xlsx(HEADERS, [
            ["ООО Импортированное 1", None, "7711111111", "771101001", None, None, None, None],
            ["ООО Импортированное 2", None, "7722222222", "772201001", None, None, None, None],
        ])
        preview = client.post(
            "/legal-entities/import",
            files={"file": ("data.xlsx", content,
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"mode": "add"},
        )
        assert preview.status_code == 200
        assert "будет добавлено — 2" in preview.text
        import re

        token = re.search(r'name="token" value="([0-9a-f]+)"', preview.text).group(1)
        confirm = client.post("/legal-entities/import/confirm", data={"token": token})
        assert confirm.status_code == 200 and "Импорт завершён" in confirm.text
        assert client.get("/legal-entities/api/list").json()["total"] == 2

    def test_scenario_4_bad_data_detected(self, client):
        """Сценарий 4: обнаружение некорректных данных при импорте."""
        content = _xlsx(HEADERS, [
            ["ООО Кривой ИНН", None, "abc123", None, None, None, None, None],
            [None, None, "7733333333", None, None, None, None, None],
        ])
        preview = client.post("/legal-entities/import",
                              files={"file": ("bad.xlsx", content,
                                              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert "отклонено — 2" in preview.text
        assert "ИНН должен содержать" in preview.text
        assert "Полное наименование" in preview.text

    def test_scenario_5_cancel_keeps_data_intact(self, client):
        """Сценарий 5: отмена импорта без повреждения данных."""
        # существующая запись до импорта
        client.post("/legal-entities/new", data={"full_name": "ООО Существующее"})
        content = _xlsx(HEADERS, [
            ["ООО Хорошее", None, "7744444444", None, None, None, None, None],
            ["ООО С плохим ОГРН", None, "7755555555", None, "999", None, None, None],
        ])
        client.post("/legal-entities/import",
                    files={"file": ("mixed.xlsx", content,
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        # «Отмена»: подтверждаем несуществующий токен — сохранение не происходит
        failed = client.post("/legal-entities/import/confirm", data={"token": "нет-такого"})
        assert failed.status_code == 422
        api = client.get("/legal-entities/api/list").json()
        assert api["total"] == 1  # только исходная запись, ничего не повреждено
        assert api["items"][0]["full_name"] == "ООО Существующее"

    def test_scenario_12_duplicate_warning(self, client):
        """Сценарий 12: предупреждение о возможном дубликате (совпадение наименования)."""
        client.post("/legal-entities/new", data={"full_name": "ООО Двойник", "inn": "7766666666"})
        content = _xlsx(HEADERS, [["ООО Двойник", None, "7777777777", None, None, None, None, None]])
        preview = client.post("/legal-entities/import",
                              files={"file": ("dup.xlsx", content,
                                              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert "возможен дубликат" in preview.text

    def test_scenario_14_export_filtered(self, client):
        """Сценарий 14: экспорт отфильтрованного списка без потерь данных."""
        client.post("/legal-entities/new", data={"full_name": "ООО Экспорт А", "inn": "7701234567"})
        client.post("/legal-entities/new", data={"full_name": "АО Экспорт Б"})
        export = client.get("/legal-entities/api/export")
        assert export.status_code == 200

        from openpyxl import load_workbook

        rows = list(load_workbook(io.BytesIO(export.content)).active.iter_rows(values_only=True))
        assert rows[0][0] == "Полное наименование"
        assert {r[0] for r in rows[1:]} == {"ООО Экспорт А", "АО Экспорт Б"}
        csv = client.get("/legal-entities/api/export?fmt=csv&q=Экспорт А")
        assert csv.content.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM (IMPORT_EXPORT §6)
        assert "ООО Экспорт А".encode() in csv.content and "АО Экспорт Б".encode() not in csv.content

    def test_archive_via_http_and_current_selection(self, client):
        resp = client.post("/legal-entities/new", data={"full_name": "ООО Архив"},
                           follow_redirects=False)
        assert resp.status_code == 303
        le_id = resp.headers["location"].rstrip("/").split("/")[-1]
        # выбор текущим (T2-07)
        sel = client.post(f"/legal-entities/{le_id}/select", follow_redirects=False)
        assert sel.status_code == 303
        assert client.cookies.get("saot_current_le") == le_id
        header = client.get("/legal-entities/")
        assert "Текущее:" in header.text
        # архивирование
        arch = client.post(f"/legal-entities/{le_id}/archive", follow_redirects=False)
        assert arch.status_code == 303
        assert client.get("/legal-entities/api/list").json()["total"] == 0
        archived = client.get("/legal-entities/?status=archived")
        assert "ООО Архив" in archived.text

    def test_dashboard_widget_zero_when_empty(self, client):
        """Главная панель: пустое состояние показывает 0 и подсказку (без фиктивных данных)."""
        home = client.get("/")
        assert "Пока нет ни одного юридического лица" in home.text
        client.post("/legal-entities/new", data={"full_name": "ООО Для панели"})
        home = client.get("/")
        assert "Юридические лица" in home.text
